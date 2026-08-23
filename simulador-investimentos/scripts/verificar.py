#!/usr/bin/env python3
"""
CONFERÊNCIA DOS NÚMEROS — testes de sanidade sobre a planilha já recalculada.

Não basta a planilha abrir sem erro de fórmula: uma referência trocada produz um
arquivo limpo, sem nenhum `#VALUE!`, e com números errados. Este script:

  1. REFAZ a projeção mês a mês em Python, partindo da regra tributária — não
     transcrevendo a fórmula da planilha. Transcrever a fórmula faria o teste
     repetir o erro junto com ela, e foi exatamente assim que um bug de
     come-cotas passou despercebido na primeira rodada.
  2. Verifica INVARIANTES que precisam valer sempre, qualquer que seja o
     produto. São eles que pegam erro de lógica sem depender de eu ter
     reimplementado a conta certa.

    python3 scripts/verificar.py [Simulador_Investimentos.xlsx]

Rode sempre depois de mexer no motor. Sai com código 1 se algo não bater.
"""

import sys
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter as L

RAIZ = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(RAIZ))
from src import refs as R          # noqa: E402

TOL = 0.005
falhas = []


def _f(v):
    if v is None:
        return "—"
    if isinstance(v, str):
        return v[:16]
    return f"{v:,.2f}"


def conf(nome, obtido, esperado, abs_tol=1.0):
    o, e = obtido or 0, esperado or 0
    ok = abs(o - e) <= max(abs_tol, abs(e) * TOL)
    print(f"  {'✓' if ok else '✗'} {nome:50s} {_f(obtido):>15s} vs {_f(esperado):>15s}")
    if not ok:
        falhas.append(nome)


def inv(nome, condicao, detalhe=""):
    print(f"  {'✓' if condicao else '✗'} {nome}" + (f"   {detalhe}" if detalhe else ""))
    if not condicao:
        falhas.append(nome)


# --------------------------------------------------------------------------
# Motor de referência, escrito a partir da REGRA, não da planilha
# --------------------------------------------------------------------------
def projetar(capital, aporte, prazo, mes_inicial, *, ret_bruto_aa, ret_bench_aa,
             taxa_adm_aa, custodia_aa, taxa_perf, carreg_ent, carreg_sai,
             tem_come_cotas, aliq_come_cotas, base_total, aliq_resgate,
             ipca_aa=0.0, aporte_corrigido=False,
             eh_pgbl=False, pgbl_reinveste=False, pgbl_beneficio=0.0):
    """Devolve (saldo_bruto, liquido, aportado_bruto, come_cotas_total, ir_resgate).

    Regra do come-cotas: em maio e novembro o fundo antecipa imposto sobre o
    ganho ACUMULADO que ainda não foi tributado. Como o ganho já tributado é
    justamente o que gerou os recolhimentos anteriores, o valor devido no
    evento é  alíquota × ganho_total_até_agora  MENOS o que já foi recolhido.

    Restituição do PGBL: quando reinvestida, volta como aporte extra em maio de
    cada ano, a partir do segundo ano (a declaração do primeiro ano só é
    entregue no ano seguinte).
    """
    i_bruto = (1 + ret_bruto_aa) ** (1 / 12) - 1
    i_bench = (1 + ret_bench_aa) ** (1 / 12) - 1
    custo_m = (taxa_adm_aa + custodia_aa) / 12

    saldo = capital * (1 - carreg_ent)
    aportado = saldo               # base fiscal (líquida de carregamento)
    aportado_bruto = capital       # o que saiu do bolso do cliente
    cc_total = 0.0

    for t in range(1, prazo + 1):
        mes_cal = (mes_inicial - 1 + t) % 12 + 1
        ap = aporte * ((1 + ipca_aa) ** (t / 12) if aporte_corrigido else 1.0)
        if eh_pgbl and pgbl_reinveste and t >= 13 and mes_cal == 5:
            ap += pgbl_beneficio
        ap_liq = ap * (1 - carreg_ent)

        rend = saldo * i_bruto
        adm = saldo * custo_m
        perf = taxa_perf * max(0.0, rend - saldo * i_bench)
        saldo = saldo + rend - adm - perf + ap_liq
        aportado += ap_liq
        aportado_bruto += ap

        if tem_come_cotas and mes_cal in (5, 11):
            ganho_total = saldo + cc_total - aportado          # ganho bruto de sempre
            devido = max(0.0, aliq_come_cotas * max(0.0, ganho_total) - cc_total)
            saldo -= devido
            cc_total += devido

    ganho_total = saldo + cc_total - aportado
    base = saldo if base_total else max(0.0, ganho_total)
    ir = max(0.0, aliq_resgate * base - cc_total)
    liquido = max(0.0, saldo - ir - saldo * carreg_sai)
    return saldo, liquido, aportado_bruto, cc_total, ir


def main():
    arq = Path(sys.argv[1]) if len(sys.argv) > 1 else RAIZ / "Simulador_Investimentos.xlsx"
    wb = load_workbook(arq, data_only=True)
    cl, mo, mi, pr, cp = (wb["Cliente"], wb["Motor"], wb["MotorImovel"],
                          wb["Premissas"], wb["Comparar"])

    cons = wb["Consultor"]
    capital, aporte, prazo = cl["F10"].value, cl["F12"].value, cl["F16"].value
    mes_ini = mo[f"B{R.MOT_GRID_INI}"].value.month
    aporte_corrigido = cons["D7"].value == "Sim"
    pgbl_reinveste = cons["D9"].value == "Sim"
    pgbl_beneficio = cons["D13"].value or 0
    print(f"\nCenário: capital {_f(capital)} · aporte {_f(aporte)}/mês · "
          f"prazo {prazo} meses · início no mês {mes_ini}")
    print(f"         aporte corrigido pela inflação: {'sim' if aporte_corrigido else 'não'} · "
          f"restituição do PGBL reinvestida: {'sim' if pgbl_reinveste else 'não'} "
          f"({_f(pgbl_beneficio)}/ano)\n")

    macro = {pr[f"B{i}"].value: pr[f"F{i}"].value
             for i in range(R.PRE_MACRO_INI, R.PRE_MACRO_FIM + 1)}
    print("Indicadores do cenário em uso:")
    for k, v in macro.items():
        print(f"    {k:10s} {v:.4%}" if isinstance(v, (int, float)) else f"    {k:10s} {v}")
    selic, ipca, tr = macro["SELIC"], macro["IPCA"], macro["TR"]

    print("\n1 · Regras de derivação da aba Premissas")
    conf("CDI = Selic − 0,10 p.p.", macro["CDI"], selic - 0.001, abs_tol=1e-6)
    conf("Poupança pela Lei 12.703",
         macro["POUP"],
         (1.005 ** 12 - 1 + tr) if selic > 0.085 else (0.7 * selic + tr), abs_tol=1e-6)

    # tabela de IR lida da planilha (é dado, não cálculo)
    regimes = {pr[f"{L(2 + j)}{R.PRE_IR_CAB}"].value: 2 + j for j in range(7)}

    def aliq_ir(regime, dias):
        col = regimes.get(regime)
        if col is None:
            return 0.0
        a = 0.0
        for i in range(R.PRE_IR_INI, R.PRE_IR_FIM + 1):
            if dias >= pr[f"A{i}"].value:
                a = pr[f"{L(col)}{i}"].value
        return a

    print("\n2 · Cada alternativa, reprojetada em Python a partir da regra")
    for k in range(R.N_SLOTS):
        c = L(2 + k)
        if mo[f"{c}{R.MOT_P['Alternativa em uso (1/0)']}"].value != 1:
            continue
        g = lambda rot: mo[f"{c}{R.MOT_P[rot]}"].value          # noqa: E731
        res = lambda rot: mo[f"{c}{R.MOT_R[rot]}"].value        # noqa: E731
        nome = g("Produto selecionado")
        regime = g("Regime tributário")
        print(f"\n  ── {nome}   [{regime}]")

        bruto, liq, aportado, cc, ir = projetar(
            capital, aporte, prazo, mes_ini,
            ret_bruto_aa=g("Retorno BRUTO a.a. (cenário)"),
            ret_bench_aa=g("Retorno a.a. do benchmark puro"),
            taxa_adm_aa=g("Taxa de administração a.a."),
            custodia_aa=g("Custódia a.a."),
            taxa_perf=g("Taxa de performance") or 0,
            carreg_ent=g("Carregamento de entrada") or 0,
            carreg_sai=g("Carregamento de saída") or 0,
            tem_come_cotas=g("Tem come-cotas (1/0)") == 1,
            aliq_come_cotas=g("Alíquota do come-cotas") or 0,
            base_total=g("Base tributável (1=rend. / 2=total)") == 2,
            aliq_resgate=aliq_ir(regime, prazo * 30),
            ipca_aa=ipca, aporte_corrigido=aporte_corrigido,
            eh_pgbl=g("É PGBL (aporte dedutível)") == 1,
            pgbl_reinveste=pgbl_reinveste, pgbl_beneficio=pgbl_beneficio)

        conf("patrimônio bruto no prazo", res("Patrimônio bruto no prazo"), bruto)
        conf("patrimônio líquido no prazo", res("Patrimônio líquido no prazo"), liq)
        conf("total investido", res("Total investido (capital + aportes)"), aportado)
        conf("imposto total (come-cotas + IR no resgate)", res("Imposto total pago"), cc + ir)
        conf("valor em R$ de hoje", res("Patrimônio líquido em R$ de hoje"),
             liq / (1 + ipca) ** (prazo / 12))

    print("\n3 · Invariantes que precisam valer sempre")
    for k in range(R.N_SLOTS):
        c = L(2 + k)
        if mo[f"{c}{R.MOT_P['Alternativa em uso (1/0)']}"].value != 1:
            continue
        res = lambda rot: mo[f"{c}{R.MOT_R[rot]}"].value or 0    # noqa: E731
        nome = mo[f"{c}{R.MOT_P['Produto selecionado']}"].value
        bruto = res("Patrimônio bruto no prazo")
        liq = res("Patrimônio líquido no prazo")
        real = res("Patrimônio líquido em R$ de hoje")
        inves = res("Total investido (capital + aportes)")
        imposto = res("Imposto total pago")
        ganho_bruto = bruto + imposto - inves
        ret_aa = mo[f"{c}{R.MOT_P['Retorno LÍQUIDO a.a. estimado']}"].value or 0

        print(f"\n  ── {nome}")
        inv("imposto nunca supera o ganho bruto", imposto <= ganho_bruto + 1,
            f"imposto {_f(imposto)} · ganho bruto {_f(ganho_bruto)}")
        inv("líquido nunca supera o bruto", liq <= bruto + 0.01,
            f"líquido {_f(liq)} · bruto {_f(bruto)}")
        inv("com inflação positiva, valor real fica abaixo do nominal",
            real <= liq + 0.01 if ipca > 0 else True)
        inv("patrimônio positivo", liq > 0, f"{_f(liq)}")
        if ret_aa > 0:
            inv("retorno líquido positivo termina acima do investido", liq > inves,
                f"líquido {_f(liq)} · investido {_f(inves)}")

    print("\n4 · Estratégia imobiliária")
    P, C = R.MIM_P, R.MIM_C
    v = lambda rot: mi[f"B{P[rot]}"].value                       # noqa: E731
    col = lambda nome, lin: mi[f"{L(C[nome])}{lin}"].value        # noqa: E731
    conf("desembolso = valor + custos de aquisição", v("Desembolso total na compra"),
         v("Valor do imóvel") + v("Custos de aquisição (R$)"), abs_tol=0.01)
    conf("sobra/falta = capital − desembolso", v("Sobra (+) ou falta (−) de caixa"),
         capital - v("Desembolso total na compra"), abs_tol=0.01)
    conf("valor do imóvel no prazo", col("Valor de mercado do imóvel", R.MIM_GRID_INI + prazo),
         v("Valor do imóvel") * (1 + v("Valorização mensal do imóvel")) ** prazo)
    conf("mês 0 não tem aluguel (a compra acabou de acontecer)",
         col("Aluguel contratado", R.MIM_GRID_INI), 0, abs_tol=0.01)
    conf("aluguel do mês 1 = aluguel inicial", col("Aluguel contratado", R.MIM_GRID_INI + 1),
         v("Aluguel mensal inicial"), abs_tol=0.01)

    lin = R.MIM_GRID_INI + prazo
    bruto_al = col("Aluguel recebido", lin)
    liq_al = col("Aluguel líquido", lin)
    inv("aluguel líquido menor que o bruto (há IPTU, taxas e imposto)",
        liq_al < bruto_al, f"líquido {_f(liq_al)} · bruto {_f(bruto_al)}")
    inv("vacância reduz o aluguel contratado",
        col("Aluguel recebido", lin) < col("Aluguel contratado", lin) + 0.01)
    if v("Sobra (+) ou falta (−) de caixa") < 0:
        print(f"    → alerta de caixa insuficiente deve estar ativo "
              f"(faltam {_f(-v('Sobra (+) ou falta (−) de caixa'))})")

    print("\n5 · Coerência entre as abas")
    for i in range(R.N_SLOTS + 1):
        origem = (mo[f"{L(2+i)}{R.MOT_R['Patrimônio líquido no prazo']}"].value
                  if i < R.N_SLOTS else mi[f"B{R.MIM_R['Patrimônio líquido no prazo']}"].value)
        destino = cp[f"C{7 + i}"].value
        rotulo = cp[f"B{7 + i}"].value
        conf(f"Comparar reflete o Motor — {str(rotulo)[:28]}", destino, origem, abs_tol=0.01)

    print("\n" + "=" * 80)
    if falhas:
        print(f"FALHOU: {len(falhas)} conferência(s) não bateram:")
        for f in dict.fromkeys(falhas):
            print(f"   · {f}")
        sys.exit(1)
    print("Todas as conferências passaram.")


if __name__ == "__main__":
    main()
