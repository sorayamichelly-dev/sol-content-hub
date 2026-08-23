#!/usr/bin/env python3
"""
EXPORTA A BASE DE PRODUTOS DA PLANILHA PARA `dados/produtos.csv`.

Fecha o ciclo de manutenção. Você cadastra e ajusta produtos no Excel — que é
onde isso é confortável de fazer —, roda este script, e a partir daí toda
planilha regerada já nasce com a sua base. Sem isso, um `build.py` sobrescreveria
seus cadastros pelos valores de referência.

    python3 scripts/exportar_produtos.py Simulador_Investimentos.xlsx

Rode depois de mexer na aba Produtos, antes de rodar build.py de novo.
"""

import csv
import sys
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError:                                   # pragma: no cover
    sys.exit("Instale a biblioteca openpyxl:  pip install openpyxl")

RAIZ = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(RAIZ))

from src import dados_seed as D           # noqa: E402
from src import refs as R                 # noqa: E402


def main():
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    origem = Path(sys.argv[1])
    if not origem.exists():
        sys.exit(f"Arquivo não encontrado: {origem}")

    wb = load_workbook(origem, data_only=False)
    if R.AB_PRODUTOS not in wb.sheetnames:
        sys.exit(f"A planilha não tem a aba {R.AB_PRODUTOS}.")
    ws = wb[R.AB_PRODUTOS]

    colunas = [c[0] for c in D.COLUNAS_PRODUTOS]
    linhas = []
    for lin in range(R.PROD_LIN_INI, R.PROD_LIN_FIM + 1):
        nome = ws.cell(row=lin, column=R.PROD_COL["nome"]).value
        if not (isinstance(nome, str) and nome.strip()):
            continue
        registro = {}
        for j, titulo in enumerate(colunas):
            v = ws.cell(row=lin, column=j + 1).value
            if isinstance(v, str) and v.startswith("="):
                print(f"  ! linha {lin}, coluna '{titulo}': há uma fórmula onde deveria "
                      f"haver um valor. Exportando em branco.")
                v = None
            registro[titulo] = "" if v is None else v
        linhas.append(registro)

    if not linhas:
        sys.exit("Nenhum produto encontrado na aba Produtos.")

    destino = RAIZ / "dados" / "produtos.csv"
    destino.parent.mkdir(exist_ok=True)
    with destino.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=colunas)
        w.writeheader()
        w.writerows(linhas)

    verificados = sum(1 for l in linhas if str(l.get("Status", "")).startswith("VERIFICADO"))
    print(f"Exportados {len(linhas)} produtos para {destino}")
    print(f"  {verificados} marcados como VERIFICADO · "
          f"{len(linhas) - verificados} ainda pendentes de conferência")
    print("Agora rode:  python3 build.py")


if __name__ == "__main__":
    main()
