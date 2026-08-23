#!/usr/bin/env python3
"""
BATERIA DE CENÁRIOS — exercita os caminhos do motor que o cenário padrão não toca.

Um simulador só está certo quando está certo para vários clientes diferentes.
Este script preenche as respostas de cinco perfis bem distintos, recalcula a
planilha em cada um e roda as conferências de `verificar.py`.

    python3 scripts/cenarios_teste.py [Simulador_Investimentos.xlsx]

Trabalha sempre sobre uma CÓPIA — o arquivo original não é tocado. (A cópia
perde os gráficos, porque openpyxl não sabe relê-los; para teste tanto faz,
e é justamente por isso que a planilha de verdade é sempre regerada por
build.py em vez de editada por script.)
"""

import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook

RAIZ = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(RAIZ))
from src import refs as R          # noqa: E402

RECALC = "/root/.claude/skills/synced/xlsx/scripts/recalc.py"

CENARIOS = [
    ("Seu João — R$ 150 mil, 10 anos, dúvida entre aplicar e comprar imóvel", {
        "Cliente": {"F10": 150000, "F12": 0, "F14": "5 a 10 anos", "F18": "Talvez",
                    "F20": "Equilíbrio entre segurança e rentabilidade", "F22": "Moderado"},
        "Consultor": {"D5": "Base", "D8": "Sim", "D12": "Não"},
        "Imovel": {"F5": "Sim", "F6": 150000, "F29": "Reinvestir"},
    }),
    ("Acumulação longa — R$ 10 mil + R$ 500/mês por 20 anos, perfil arrojado", {
        "Cliente": {"F10": 10000, "F12": 500, "F14": "Prazo personalizado", "F15": 240,
                    "F18": "Não", "F20": "Rentabilidade, aceitando oscilação",
                    "F22": "Arrojado"},
        "Consultor": {"D5": "Otimista", "D7": "Sim", "D8": "Sim", "D12": "Não"},
        "Imovel": {"F5": "Não"},
    }),
    ("Reserva de curto prazo — R$ 50 mil por 12 meses, pode precisar sacar", {
        "Cliente": {"F10": 50000, "F12": 0, "F14": "Até 1 ano", "F18": "Sim",
                    "F20": "Poder retirar a qualquer momento", "F22": "Conservador"},
        "Consultor": {"D5": "Conservador", "D8": "Sim", "D12": "Não"},
        "Imovel": {"F5": "Não"},
    }),
    ("Imóvel viável — R$ 200 mil disponíveis, casa de R$ 150 mil, aluguel consumido", {
        "Cliente": {"F10": 200000, "F12": 0, "F14": "5 a 10 anos", "F18": "Não",
                    "F20": "Segurança acima de tudo", "F22": "Conservador"},
        "Consultor": {"D5": "Base", "D8": "Sim", "D12": "Não"},
        "Imovel": {"F5": "Sim", "F6": 150000, "F13": 1100,
                   "F29": "Consumir (renda mensal)", "F26": "Sim"},
    }),
    ("PGBL com benefício fiscal reinvestido — renda R$ 120 mil, declara completo", {
        "Cliente": {"F10": 30000, "F12": 1000, "F14": "Prazo personalizado", "F15": 240,
                    "F18": "Não", "F20": "Equilíbrio entre segurança e rentabilidade",
                    "F22": "Moderado"},
        "Consultor": {"D5": "Base", "D8": "Não", "D9": "Sim", "D10": 120000,
                      "D11": 0.275, "D12": "Sim",
                      "C17": "PGBL CAIXA RF — Regressivo",
                      "C18": "VGBL CAIXA RF — Regressivo",
                      "C19": "LCI CAIXA 12 meses",
                      "C20": "Tesouro IPCA+ 2035",
                      "C21": "Fundo CAIXA Ações Ibovespa",
                      "C22": "Poupança"},
        "Imovel": {"F5": "Não"},
    }),
]


def aplicar(arquivo, mudancas):
    wb = load_workbook(arquivo)
    for aba, celulas in mudancas.items():
        ws = wb[aba]
        for ref, valor in celulas.items():
            ws[ref] = valor
    wb.save(arquivo)


def main():
    origem = Path(sys.argv[1]) if len(sys.argv) > 1 else RAIZ / "Simulador_Investimentos.xlsx"
    if not origem.exists():
        sys.exit(f"Gere a planilha primeiro:  python3 build.py\n(não encontrei {origem})")

    reprovados = []
    with tempfile.TemporaryDirectory(prefix="cenarios_") as tmp:
        for i, (nome, mudancas) in enumerate(CENARIOS, 1):
            print("\n" + "=" * 80)
            print(f"CENÁRIO {i}/{len(CENARIOS)} · {nome}")
            print("=" * 80)

            copia = Path(tmp) / f"c{i}.xlsx"
            shutil.copy(origem, copia)
            aplicar(copia, mudancas)

            r = subprocess.run([sys.executable, RECALC, str(copia), "480"],
                               capture_output=True, text=True)
            if '"status": "success"' not in r.stdout:
                print(r.stdout.strip() or r.stderr.strip())
                reprovados.append(f"{nome} (recálculo)")
                continue
            print(r.stdout.strip())

            r = subprocess.run([sys.executable, str(RAIZ / "scripts" / "verificar.py"),
                                str(copia)], capture_output=True, text=True)
            print(r.stdout.strip()[-3000:])
            if r.returncode != 0:
                reprovados.append(f"{nome} (conferências)")

    print("\n" + "=" * 80)
    if reprovados:
        print(f"REPROVOU em {len(reprovados)} cenário(s):")
        for c in reprovados:
            print(f"   · {c}")
        sys.exit(1)
    print(f"Os {len(CENARIOS)} cenários passaram em todas as conferências.")


if __name__ == "__main__":
    main()
