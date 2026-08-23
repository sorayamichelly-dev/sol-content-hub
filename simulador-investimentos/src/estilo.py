"""
Paleta, estilos e utilitários de formatação do Simulador de Investimentos.

Camada de APRESENTAÇÃO. Não contém regra de negócio nem dado de produto.
Alterar cores/fontes aqui muda a planilha inteira sem tocar em fórmula alguma.
"""

from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# --------------------------------------------------------------------------
# PALETA (corporativa, sóbria — azul/laranja)
# --------------------------------------------------------------------------
AZUL_ESCURO = "0B3B60"
AZUL = "005CA9"
AZUL_MEDIO = "2E7EBB"
AZUL_CLARO = "E8F1F9"
AZUL_MUITO_CLARO = "F4F9FD"
LARANJA = "E8730C"
LARANJA_CLARO = "FDF1E3"
CINZA_TEXTO = "3C3C3C"
CINZA_SUAVE = "7A7A7A"
CINZA_LINHA = "D6DCE4"
CINZA_FUNDO = "F5F7FA"
BRANCO = "FFFFFF"
VERDE = "1F7A3D"
VERDE_CLARO = "E6F4EA"
VERMELHO = "B3261E"
VERMELHO_CLARO = "FCEAE8"
AMBAR = "8A6100"
AMBAR_CLARO = "FFF6DF"

# Convenção de modelagem financeira (ver README):
COR_INPUT = "0000FF"       # azul: valor digitado pelo usuário
COR_FORMULA = "000000"     # preto: cálculo na própria aba
COR_LINK = "008000"        # verde: referência a outra aba
FILL_INPUT = "FFF2CC"      # amarelo claro: célula que o usuário deve preencher

FONTE = "Arial"

# --------------------------------------------------------------------------
# FORMATOS NUMÉRICOS
# --------------------------------------------------------------------------
FMT_MOEDA = '"R$" #,##0;("R$" #,##0);"—"'
FMT_MOEDA_C = '"R$" #,##0.00;("R$" #,##0.00);"—"'
FMT_NUM = '#,##0;(#,##0);"—"'
FMT_NUM2 = '#,##0.00;(#,##0.00);"—"'
FMT_PCT = '0.0%;(0.0%);"—"'
FMT_PCT2 = '0.00%;(0.00%);"—"'
FMT_DATA = "dd/mm/yyyy"
FMT_TEXTO = "@"
FMT_INT = '#,##0;(#,##0);"—"'

# --------------------------------------------------------------------------
# BORDAS
# --------------------------------------------------------------------------
_fina = Side(style="thin", color=CINZA_LINHA)
_media = Side(style="medium", color=AZUL)
BORDA_GRADE = Border(left=_fina, right=_fina, top=_fina, bottom=_fina)
BORDA_BASE = Border(bottom=Side(style="thin", color=CINZA_LINHA))
BORDA_TOPO_AZUL = Border(top=_media)
BORDA_CAIXA = Border(
    left=Side(style="thin", color=AZUL_MEDIO),
    right=Side(style="thin", color=AZUL_MEDIO),
    top=Side(style="thin", color=AZUL_MEDIO),
    bottom=Side(style="thin", color=AZUL_MEDIO),
)


def fill(cor):
    return PatternFill("solid", fgColor=cor)


def fonte(tam=10, negrito=False, cor=CINZA_TEXTO, italico=False, sublinhado=None):
    return Font(name=FONTE, size=tam, bold=negrito, color=cor,
                italic=italico, underline=sublinhado)


AL_ESQ = Alignment(horizontal="left", vertical="center", wrap_text=False)
AL_ESQ_WRAP = Alignment(horizontal="left", vertical="top", wrap_text=True)
AL_CENTRO = Alignment(horizontal="center", vertical="center", wrap_text=False)
AL_CENTRO_WRAP = Alignment(horizontal="center", vertical="center", wrap_text=True)
AL_DIR = Alignment(horizontal="right", vertical="center")


# --------------------------------------------------------------------------
# HELPERS DE ESCRITA
# --------------------------------------------------------------------------
def escreve(ws, ref, valor, *, tam=10, negrito=False, cor=CINZA_TEXTO,
            bg=None, fmt=None, alinha=None, borda=None, italico=False,
            wrap=False, sublinhado=None):
    """Escreve uma célula já formatada e devolve o objeto célula."""
    c = ws[ref]
    c.value = valor
    c.font = fonte(tam, negrito, cor, italico, sublinhado)
    if bg:
        c.fill = fill(bg)
    if fmt:
        c.number_format = fmt
    if alinha is not None:
        c.alignment = alinha
    elif wrap:
        c.alignment = AL_ESQ_WRAP
    if borda:
        c.border = borda
    return c


def titulo_pagina(ws, texto, subtitulo=None, ate_col="N"):
    """Faixa de título no topo da aba."""
    ws.merge_cells(f"A1:{ate_col}1")
    escreve(ws, "A1", f"   {texto}", tam=16, negrito=True, cor=BRANCO,
            bg=AZUL_ESCURO, alinha=AL_ESQ)
    ws.row_dimensions[1].height = 34
    if subtitulo:
        ws.merge_cells(f"A2:{ate_col}2")
        escreve(ws, "A2", f"   {subtitulo}", tam=9, cor=BRANCO,
                bg=AZUL, alinha=AL_ESQ)
        ws.row_dimensions[2].height = 18


def secao(ws, linha, texto, ate_col="N", cor_bg=AZUL_CLARO, cor_txt=AZUL_ESCURO):
    """Faixa de seção dentro da aba."""
    ws.merge_cells(f"A{linha}:{ate_col}{linha}")
    escreve(ws, f"A{linha}", f"  {texto}", tam=11, negrito=True,
            cor=cor_txt, bg=cor_bg, alinha=AL_ESQ)
    ws.row_dimensions[linha].height = 22


def nota(ws, linha, texto, ate_col="N", cor=CINZA_SUAVE, bg=None, tam=8.5,
         altura=None, negrito=False):
    ws.merge_cells(f"A{linha}:{ate_col}{linha}")
    # Uma fórmula não pode receber o recuo de dois espaços: ela viraria texto.
    conteudo = texto if isinstance(texto, str) and texto.startswith("=") else f"  {texto}"
    escreve(ws, f"A{linha}", conteudo, tam=tam, cor=cor, bg=bg,
            alinha=AL_ESQ_WRAP, negrito=negrito)
    if altura:
        ws.row_dimensions[linha].height = altura


def aviso(ws, linha, texto, ate_col="N"):
    """Caixa de alerta âmbar."""
    ws.merge_cells(f"A{linha}:{ate_col}{linha}")
    escreve(ws, f"A{linha}", f"  ⚠  {texto}", tam=9, negrito=True,
            cor=AMBAR, bg=AMBAR_CLARO, alinha=AL_ESQ_WRAP)
    ws.row_dimensions[linha].height = 26


def cabecalho_tabela(ws, linha, col_ini, valores, larguras=None, altura=30):
    """Linha de cabeçalho de tabela."""
    for i, v in enumerate(valores):
        col = get_column_letter(col_ini + i)
        escreve(ws, f"{col}{linha}", v, tam=8.5, negrito=True, cor=BRANCO,
                bg=AZUL, alinha=AL_CENTRO_WRAP, borda=BORDA_GRADE)
        if larguras:
            ws.column_dimensions[col].width = larguras[i]
    ws.row_dimensions[linha].height = altura


def celula_input(ws, ref, valor, fmt=None, tam=11, alinha=None):
    """Célula que o usuário preenche: fundo amarelo, texto azul."""
    return escreve(ws, ref, valor, tam=tam, negrito=True, cor=COR_INPUT,
                   bg=FILL_INPUT, fmt=fmt, borda=BORDA_CAIXA,
                   alinha=alinha or AL_ESQ)


def celula_resultado(ws, ref, valor, fmt=FMT_MOEDA, tam=14, cor=AZUL_ESCURO,
                     bg=AZUL_MUITO_CLARO):
    return escreve(ws, ref, valor, tam=tam, negrito=True, cor=cor, bg=bg,
                   fmt=fmt, alinha=AL_CENTRO, borda=BORDA_CAIXA)


def link_aba(ws, ref, aba, rotulo, destino="A1", bg=AZUL, cor=BRANCO, tam=10):
    """Botão de navegação (hyperlink interno) — substitui macro de navegação."""
    from openpyxl.worksheet.hyperlink import Hyperlink
    c = escreve(ws, ref, rotulo, tam=tam, negrito=True, cor=cor, bg=bg,
                alinha=AL_CENTRO, borda=BORDA_CAIXA)
    c.hyperlink = Hyperlink(ref=ref, location=f"{aba}!{destino}",
                            tooltip=f"Ir para a aba {aba}")
    return c


def larguras(ws, mapa):
    for col, larg in mapa.items():
        ws.column_dimensions[col].width = larg


def congela(ws, ref):
    ws.freeze_panes = ref


def esconde_grade(ws):
    ws.sheet_view.showGridLines = False


def define_nome(wb, nome, referencia):
    """Cria/atualiza um nome definido (named range)."""
    from openpyxl.workbook.defined_name import DefinedName
    if nome in wb.defined_names:
        del wb.defined_names[nome]
    wb.defined_names.add(DefinedName(nome, attr_text=referencia))
