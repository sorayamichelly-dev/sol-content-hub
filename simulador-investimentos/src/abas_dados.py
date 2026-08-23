"""
CAMADA DE DADOS — abas Produtos, Premissas, Fontes e Glossário.

Nenhuma fórmula do motor mora aqui. Estas abas são o "banco de dados" da
ferramenta: alterar uma taxa, cadastrar um produto novo ou atualizar um
indicador não exige tocar em cálculo nenhum.
"""

from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.comments import Comment

from .estilo import *          # noqa: F403
from .refs import *            # noqa: F403
from . import dados_seed as D


def _dv(ws, formula1, celulas, titulo, mensagem):
    dv = DataValidation(type="list", formula1=formula1, allow_blank=True,
                        showDropDown=False, showErrorMessage=True)
    dv.promptTitle = titulo
    dv.prompt = mensagem
    dv.showInputMessage = True
    ws.add_data_validation(dv)
    for c in celulas:
        dv.add(c)
    return dv


def _lista(valores):
    return '"' + ",".join(valores) + '"'


# ==========================================================================
# ABA PRODUTOS
# ==========================================================================
def construir_produtos(wb):
    ws = wb[AB_PRODUTOS]
    esconde_grade(ws)
    titulo_pagina(ws, "BASE DE PRODUTOS",
                  "Camada de dados — altere aqui sem mexer em nenhuma fórmula",
                  ate_col="AD")

    secao(ws, 4, "COMO USAR ESTA BASE", ate_col="AD")
    nota(ws, 5,
         "Cada linha é um produto. Para cadastrar um novo, copie a linha ZZ99 (a última, marcada como EXEMPLO), "
         "cole em uma linha vazia abaixo e ajuste os campos. O produto passa a aparecer automaticamente nas listas "
         "de escolha da aba Consultor. Preencha a coluna Ativo com SIM para colocá-lo em uso.",
         ate_col="AD", altura=30)
    nota(ws, 6,
         "Legenda de cores da planilha inteira:  texto AZUL = valor digitado por você  ·  texto PRETO = cálculo da própria aba  "
         "·  texto VERDE = valor que vem de outra aba  ·  fundo AMARELO = célula que você deve preencher.",
         ate_col="AD", altura=22)
    aviso(ws, 7,
          "As taxas abaixo nascem como REFERÊNCIA. Antes de usar em atendimento, confirme produto a produto na fonte "
          "oficial (lâmina, regulamento, tabela de taxas vigente) e mude a coluna Status para VERIFICADO.",
          ate_col="AD")

    # Cabeçalho
    larg = [c[1] for c in D.COLUNAS_PRODUTOS]
    cabecalho_tabela(ws, PROD_LIN_CAB, 1, [c[0] for c in D.COLUNAS_PRODUTOS],
                     larguras=larg, altura=34)

    dicas = {
        "Benchmark": "Índice de referência. Deve existir na aba Premissas (coluna Código).",
        "% do benchmark": "Quanto do índice o produto entrega, BRUTO, antes da taxa de administração. 1,00 = 100%.",
        "Spread a.a.": "Ganho adicional sobre o índice. Em IPCA+6%, digite 6% aqui.",
        "Composição": "SOMA: taxa = índice × % + spread.  MULT: taxa = (1 + índice × %) × (1 + spread) − 1 (use em IPCA+).",
        "Base tributável": "1 = IR só sobre o rendimento (regra geral e VGBL).  2 = IR sobre o valor total resgatado (PGBL).",
        "Carência (dias)": "Tempo mínimo antes de poder resgatar. 0 = liquidez imediata.",
        "Liquidez (dias)": "Depois da carência, quantos dias o dinheiro leva para cair na conta.",
        "Regime tributário": "Ver a aba Premissas: ISENTO, RF_REGR, FUNDO_CP, FUNDO_LP, FUNDO_ACOES, PREV_REGR, PREV_PROG, IMOVEL.",
        "Risco (1-5)": "1 = muito baixo … 5 = muito alto. Alimenta o filtro de perfil e o texto mostrado ao cliente.",
        "Rent. 12m (obs.)": "DADO OBSERVADO (passado). Nunca é usado na projeção — serve só para conversa e contexto.",
        "Tipo de dado": "Tempo real / Diário (D+1) / Mensal / Histórico / Cadastral / Premissa / Calculado.",
    }
    for nome_col, dica in dicas.items():
        idx = [c[0] for c in D.COLUNAS_PRODUTOS].index(nome_col) + 1
        ws.cell(row=PROD_LIN_CAB, column=idx).comment = Comment(dica, "Simulador")

    # Dados
    for i, prod in enumerate(D.PRODUTOS):
        lin = PROD_LIN_INI + i
        exemplo = prod[0] == "ZZ99"
        for j, valor in enumerate(prod):
            col = get_column_letter(j + 1)
            nome_campo = D.COLUNAS_PRODUTOS[j][0]
            fmt, alin = None, AL_ESQ
            if nome_campo in ("% do benchmark",):
                fmt, alin = FMT_PCT, AL_CENTRO
            elif nome_campo in ("Spread a.a.", "Taxa adm. a.a.", "Taxa perf. (%)",
                                "Carreg. entrada", "Carreg. saída", "Custódia a.a.",
                                "Volatilidade a.a.", "Rent. 12m (obs.)",
                                "Rent. 36m a.a. (obs.)"):
                fmt, alin = FMT_PCT2, AL_CENTRO
            elif nome_campo in ("Aplicação mínima", "Aporte mín. mensal"):
                fmt, alin = FMT_MOEDA, AL_DIR
            elif nome_campo in ("Carência (dias)", "Liquidez (dias)", "Risco (1-5)",
                                "Base tributável"):
                fmt, alin = FMT_INT, AL_CENTRO
            elif nome_campo == "Atualizado em":
                fmt, alin = FMT_TEXTO, AL_CENTRO
            elif nome_campo in ("Categoria", "Benchmark", "Composição",
                                "Regime tributário", "Ativo", "Status", "Tipo de dado"):
                alin = AL_CENTRO

            escreve(ws, f"{col}{lin}", valor, tam=8.5,
                    cor=CINZA_SUAVE if exemplo else COR_INPUT,
                    italico=exemplo, fmt=fmt, alinha=alin, borda=BORDA_GRADE,
                    bg=CINZA_FUNDO if exemplo else None)
        ws.row_dimensions[lin].height = 14

    # Linhas em branco, já formatadas, prontas para novos cadastros
    for lin in range(PROD_LIN_INI + len(D.PRODUTOS), PROD_LIN_FIM + 1):
        for j in range(len(D.COLUNAS_PRODUTOS)):
            ws.cell(row=lin, column=j + 1).border = BORDA_GRADE
        ws.row_dimensions[lin].height = 14

    # Validações
    faixa = f"A{PROD_LIN_INI}:AD{PROD_LIN_FIM}"
    fim = PROD_LIN_FIM
    _dv(ws, _lista(["Renda Fixa", "Fundo", "Previdência", "Imobiliário", "Outros"]),
        [f"C{PROD_LIN_INI}:C{fim}"], "Categoria", "Escolha a categoria do produto.")
    _dv(ws, f"{AB_PREMISSAS}!$B${PRE_MACRO_INI}:$B${PRE_MACRO_FIM}",
        [f"E{PROD_LIN_INI}:E{fim}"], "Benchmark", "Deve existir na aba Premissas.")
    _dv(ws, _lista(["SOMA", "MULT"]), [f"H{PROD_LIN_INI}:H{fim}"],
        "Composição", "MULT para produtos IPCA+; SOMA para os demais.")
    _dv(ws, _lista(D.REGIMES + ["IMOVEL"]), [f"N{PROD_LIN_INI}:N{fim}"],
        "Regime tributário", "Define como o IR é calculado.")
    _dv(ws, _lista(["1", "2"]), [f"O{PROD_LIN_INI}:O{fim}"],
        "Base tributável", "1 = só o rendimento · 2 = o valor total (PGBL).")
    _dv(ws, _lista(["1", "2", "3", "4", "5"]), [f"T{PROD_LIN_INI}:T{fim}"],
        "Risco", "1 = muito baixo … 5 = muito alto.")
    _dv(ws, _lista(["Tempo real", "Diário (D+1)", "Mensal", "Histórico",
                    "Cadastral", "Premissa", "Calculado"]),
        [f"AA{PROD_LIN_INI}:AA{fim}"], "Tipo de dado",
        "Diga de que natureza é o dado — evita tratar dado velho como tempo real.")
    _dv(ws, _lista(["VERIFICADO", "EXEMPLO — VERIFICAR", "DESATUALIZADO"]),
        [f"AB{PROD_LIN_INI}:AB{fim}"], "Status", "Controle de confiabilidade do cadastro.")
    _dv(ws, _lista(["SIM", "NÃO"]), [f"AD{PROD_LIN_INI}:AD{fim}"],
        "Ativo", "SIM coloca o produto em uso nas simulações.")

    # Realce dos produtos ainda não conferidos
    ws.conditional_formatting.add(
        faixa,
        FormulaRule(formula=[f'$AB{PROD_LIN_INI}="EXEMPLO — VERIFICAR"'],
                    fill=fill(LARANJA_CLARO), stopIfTrue=False))
    ws.conditional_formatting.add(
        faixa,
        FormulaRule(formula=[f'$AB{PROD_LIN_INI}="DESATUALIZADO"'],
                    fill=fill(VERMELHO_CLARO), stopIfTrue=False))

    ws.freeze_panes = f"C{PROD_LIN_INI}"
    ws.auto_filter.ref = f"A{PROD_LIN_CAB}:AD{PROD_LIN_FIM}"
    ws.sheet_properties.tabColor = LARANJA
    link_aba(ws, "AF1", AB_INICIO, "◀ INÍCIO")


# ==========================================================================
# ABA PREMISSAS
# ==========================================================================
def construir_premissas(wb):
    ws = wb[AB_PREMISSAS]
    esconde_grade(ws)
    titulo_pagina(ws, "PREMISSAS E PARÂMETROS DE CÁLCULO",
                  "Camada de dados — tudo que o motor usa está aqui, nada está escondido em fórmula",
                  ate_col="J")
    larguras(ws, {"A": 42, "B": 15, "C": 15, "D": 15, "E": 15, "F": 17,
                  "G": 13, "H": 40, "I": 22, "J": 30})

    escreve(ws, "A4", "Data de atualização dos dados", tam=10, negrito=True)
    celula_input(ws, "B4", D.DATA_REFERENCIA, tam=10, alinha=AL_CENTRO)
    escreve(ws, "A5", "Atualizado por", tam=10, negrito=True)
    celula_input(ws, "B5", "(preencher)", tam=10, alinha=AL_CENTRO)
    escreve(ws, "A6", "Cenário selecionado na aba Consultor", tam=10, negrito=True)
    escreve(ws, "B6", "=CENARIO", tam=10, negrito=True, cor=COR_LINK,
            alinha=AL_CENTRO, bg=VERDE_CLARO, borda=BORDA_CAIXA)

    # ---------------- BLOCO 1: MACRO ----------------
    secao(ws, 7 - 1 + 0, "1 · INDICADORES MACROECONÔMICOS  (% ao ano)", ate_col="J")
    cab = ["Indicador", "Código", "Conservador", "Base", "Otimista",
           "No cenário em uso", "Composição", "Fonte", "Tipo de dado", "Observação"]
    cabecalho_tabela(ws, PRE_MACRO_CAB, 1, cab, altura=30)

    for i, (cod, nome, cons, base, oti, comp, fonte_txt, tipo) in enumerate(D.MACRO):
        lin = PRE_MACRO_INI + i
        escreve(ws, f"A{lin}", nome, tam=9, borda=BORDA_GRADE)
        escreve(ws, f"B{lin}", cod, tam=9, negrito=True, alinha=AL_CENTRO,
                borda=BORDA_GRADE, cor=AZUL_ESCURO)
        for col, val in zip(("C", "D", "E"), (cons, base, oti)):
            if val is None:
                ws[f"{col}{lin}"].value = None      # preenchido por fórmula abaixo
                escreve(ws, f"{col}{lin}", None, tam=9, fmt=FMT_PCT2,
                        alinha=AL_CENTRO, borda=BORDA_GRADE)
            else:
                escreve(ws, f"{col}{lin}", val, tam=9, cor=COR_INPUT, negrito=True,
                        bg=FILL_INPUT, fmt=FMT_PCT2, alinha=AL_CENTRO, borda=BORDA_GRADE)
        escreve(ws, f"F{lin}",
                f"=INDEX(C{lin}:E{lin},MATCH(CENARIO,$C${PRE_MACRO_CAB}:$E${PRE_MACRO_CAB},0))",
                tam=9, negrito=True, cor=AZUL_ESCURO, bg=AZUL_CLARO,
                fmt=FMT_PCT2, alinha=AL_CENTRO, borda=BORDA_GRADE)
        escreve(ws, f"G{lin}", comp, tam=9, alinha=AL_CENTRO, borda=BORDA_GRADE)
        escreve(ws, f"H{lin}", fonte_txt, tam=8, cor=CINZA_SUAVE, borda=BORDA_GRADE)
        escreve(ws, f"I{lin}", tipo, tam=8, cor=CINZA_SUAVE, alinha=AL_CENTRO,
                borda=BORDA_GRADE)
        escreve(ws, f"J{lin}", "", tam=8, borda=BORDA_GRADE)

    lin_selic = PRE_MACRO_INI
    lin_cdi = PRE_MACRO_INI + 1
    lin_tr = PRE_MACRO_INI + 3
    lin_poup = PRE_MACRO_INI + 4
    for col in ("C", "D", "E"):
        escreve(ws, f"{col}{lin_cdi}", f"={col}{lin_selic}-0.001", tam=9,
                fmt=FMT_PCT2, alinha=AL_CENTRO, borda=BORDA_GRADE)
        escreve(ws, f"{col}{lin_poup}",
                f"=IF({col}{lin_selic}>0.085,1.005^12-1+{col}{lin_tr},"
                f"0.7*{col}{lin_selic}+{col}{lin_tr})",
                tam=9, fmt=FMT_PCT2, alinha=AL_CENTRO, borda=BORDA_GRADE)
    escreve(ws, f"J{lin_cdi}", "Calculado: Selic − 0,10 p.p.", tam=8, cor=CINZA_SUAVE,
            borda=BORDA_GRADE)
    escreve(ws, f"J{lin_poup}",
            "Calculado: se Selic > 8,5% a.a. → 0,5% a.m. + TR; senão 70% da Selic + TR",
            tam=8, cor=CINZA_SUAVE, borda=BORDA_GRADE)

    nota(ws, PRE_MACRO_FIM + 1,
         "Cenário NÃO é promessa de rentabilidade. É um conjunto de premissas escolhido pelo consultor "
         "para mostrar ao cliente uma faixa de resultados possíveis.", ate_col="J", altura=18)

    # ---------------- BLOCO 2: IR ----------------
    secao(ws, PRE_IR_CAB - 2, "2 · IMPOSTO DE RENDA NO RESGATE  (alíquota por regime e prazo)", ate_col="J")
    nota(ws, PRE_IR_CAB - 1,
         "Leitura: a linha usada é a maior faixa de dias que o prazo já alcançou. "
         "Ex.: 400 dias em RF_REGR cai na faixa de 361 dias → 17,5%.", ate_col="J")
    cabecalho_tabela(ws, PRE_IR_CAB, 1, ["Dias mínimos"] + D.REGIMES, altura=26)
    for i, linha_ir in enumerate(D.TABELA_IR):
        lin = PRE_IR_INI + i
        escreve(ws, f"A{lin}", linha_ir[0], tam=9, negrito=True, fmt=FMT_INT,
                alinha=AL_CENTRO, borda=BORDA_GRADE, cor=AZUL_ESCURO)
        for j, aliq in enumerate(linha_ir[1:]):
            escreve(ws, f"{get_column_letter(2 + j)}{lin}", aliq, tam=9,
                    cor=COR_INPUT, bg=FILL_INPUT, fmt=FMT_PCT, alinha=AL_CENTRO,
                    borda=BORDA_GRADE)
    for j, reg in enumerate(D.REGIMES):
        ws.cell(row=PRE_IR_CAB, column=2 + j).comment = Comment(
            D.REGIME_DESCRICAO[reg], "Simulador")

    # ---------------- BLOCO 3: COME-COTAS ----------------
    secao(ws, PRE_CC_CAB - 1, "3 · COME-COTAS  (antecipação semestral de IR — maio e novembro)", ate_col="J")
    cabecalho_tabela(ws, PRE_CC_CAB, 1,
                     ["Regime", "Tem come-cotas (1/0)", "Alíquota", "", "",
                      "", "", "Como funciona", "", ""], altura=26)
    for i, (reg, tem, aliq) in enumerate(D.COME_COTAS):
        lin = PRE_CC_INI + i
        escreve(ws, f"A{lin}", reg, tam=9, negrito=True, borda=BORDA_GRADE)
        escreve(ws, f"B{lin}", tem, tam=9, cor=COR_INPUT, bg=FILL_INPUT,
                fmt=FMT_INT, alinha=AL_CENTRO, borda=BORDA_GRADE)
        escreve(ws, f"C{lin}", aliq, tam=9, cor=COR_INPUT, bg=FILL_INPUT,
                fmt=FMT_PCT, alinha=AL_CENTRO, borda=BORDA_GRADE)
        escreve(ws, f"H{lin}", D.REGIME_DESCRICAO[reg], tam=8, cor=CINZA_SUAVE,
                borda=BORDA_GRADE)

    # ---------------- BLOCO 4: IOF ----------------
    secao(ws, PRE_IOF_CAB - 1, "4 · IOF REGRESSIVO  (% do rendimento, resgates com menos de 30 dias)", ate_col="J")
    cabecalho_tabela(ws, PRE_IOF_CAB, 1,
                     ["Dias de aplicação", "IOF sobre o rendimento", "", "", "",
                      "", "", "", "", ""], altura=22)
    for i, pct in enumerate(D.TABELA_IOF):
        lin = PRE_IOF_INI + i
        escreve(ws, f"A{lin}", i + 1, tam=8.5, fmt=FMT_INT, alinha=AL_CENTRO,
                borda=BORDA_GRADE)
        escreve(ws, f"B{lin}", pct, tam=8.5, cor=COR_INPUT, bg=FILL_INPUT,
                fmt=FMT_PCT, alinha=AL_CENTRO, borda=BORDA_GRADE)
        ws.row_dimensions[lin].height = 12
    escreve(ws, f"H{PRE_IOF_INI}",
            "A projeção mês a mês trata cada mês como 30 dias, então o IOF só aparece "
            "em simulações de menos de um mês. Para resgates de poucos dias use a "
            "calculadora de curto prazo na aba Consultor.",
            tam=8, cor=CINZA_SUAVE, alinha=AL_ESQ_WRAP)

    # ---------------- BLOCO 5: IRPF MENSAL ----------------
    secao(ws, PRE_IRPF_CAB - 1,
          "5 · TABELA PROGRESSIVA MENSAL DO IRPF  (carnê-leão do aluguel)", ate_col="J")
    cabecalho_tabela(ws, PRE_IRPF_CAB, 1,
                     ["Base mínima da faixa (R$/mês)", "Alíquota", "Parcela a deduzir (R$)",
                      "", "", "", "", "Fonte", "", ""], altura=26)
    for i, (base, aliq, ded) in enumerate(D.TABELA_IRPF_MENSAL):
        lin = PRE_IRPF_INI + i
        escreve(ws, f"A{lin}", base, tam=9, cor=COR_INPUT, bg=FILL_INPUT,
                fmt=FMT_MOEDA_C, alinha=AL_CENTRO, borda=BORDA_GRADE)
        escreve(ws, f"B{lin}", aliq, tam=9, cor=COR_INPUT, bg=FILL_INPUT,
                fmt=FMT_PCT, alinha=AL_CENTRO, borda=BORDA_GRADE)
        escreve(ws, f"C{lin}", ded, tam=9, cor=COR_INPUT, bg=FILL_INPUT,
                fmt=FMT_MOEDA_C, alinha=AL_CENTRO, borda=BORDA_GRADE)
    escreve(ws, f"H{PRE_IRPF_INI}", "Receita Federal — conferir a vigência a cada exercício.",
            tam=8, cor=CINZA_SUAVE)

    # ---------------- BLOCO 6: HORIZONTES ----------------
    secao(ws, PRE_HOR_CAB - 1, "6 · PRAZOS OFERECIDOS AO CLIENTE", ate_col="J")
    cabecalho_tabela(ws, PRE_HOR_CAB, 1,
                     ["Como o cliente responde", "Meses usados no cálculo", "", "",
                      "", "", "", "", "", ""], altura=22)
    for i, (rotulo, meses) in enumerate(D.HORIZONTES):
        lin = PRE_HOR_INI + i
        escreve(ws, f"A{lin}", rotulo, tam=9, borda=BORDA_GRADE)
        escreve(ws, f"B{lin}", meses, tam=9, cor=COR_INPUT, bg=FILL_INPUT,
                fmt=FMT_INT, alinha=AL_CENTRO, borda=BORDA_GRADE)

    # ---------------- BLOCO 7: RISCO EM LINGUAGEM SIMPLES ----------------
    secao(ws, PRE_RISCO_CAB - 1, "7 · COMO O RISCO É EXPLICADO AO CLIENTE", ate_col="J")
    cabecalho_tabela(ws, PRE_RISCO_CAB, 1,
                     ["Nível (1-5)", "Como aparece para o cliente", "Frase de apoio",
                      "", "", "", "", "", "", ""], altura=22)
    for nivel, (rotulo, frase) in D.ESCALA_RISCO.items():
        lin = PRE_RISCO_INI + nivel - 1
        escreve(ws, f"A{lin}", nivel, tam=9, negrito=True, fmt=FMT_INT,
                alinha=AL_CENTRO, borda=BORDA_GRADE, cor=AZUL_ESCURO)
        escreve(ws, f"B{lin}", rotulo, tam=9, cor=COR_INPUT, bg=FILL_INPUT,
                alinha=AL_CENTRO, borda=BORDA_GRADE)
        escreve(ws, f"C{lin}", frase, tam=8.5, cor=COR_INPUT, bg=FILL_INPUT,
                borda=BORDA_GRADE)

    # ---------------- BLOCO 8: PERFIL ----------------
    secao(ws, PRE_PERFIL_CAB - 1, "8 · PERFIL DO INVESTIDOR → RISCO MÁXIMO ACEITO", ate_col="J")
    cabecalho_tabela(ws, PRE_PERFIL_CAB, 1,
                     ["Perfil (suitability)", "Risco máximo (1-5)", "", "", "",
                      "", "", "Observação", "", ""], altura=22)
    for i, perfil in enumerate(D.PERFIS):
        lin = PRE_PERFIL_INI + i
        escreve(ws, f"A{lin}", perfil, tam=9, borda=BORDA_GRADE)
        escreve(ws, f"B{lin}", D.TETO_RISCO[perfil], tam=9, cor=COR_INPUT,
                bg=FILL_INPUT, fmt=FMT_INT, alinha=AL_CENTRO, borda=BORDA_GRADE)
    escreve(ws, f"H{PRE_PERFIL_INI}",
            "O perfil oficial vem do questionário de suitability do sistema da instituição. "
            "Aqui ele só serve para filtrar o que a planilha sugere.",
            tam=8, cor=CINZA_SUAVE, alinha=AL_ESQ_WRAP)

    # ---------------- BLOCO 9: OUTROS PARÂMETROS ----------------
    secao(ws, PRE_OUTROS_INI - 1, "9 · OUTROS PARÂMETROS", ate_col="J")
    outros = [
        ("Limite do FGC por CPF e por instituição", 250000.0, FMT_MOEDA,
         "Fundo Garantidor de Créditos — usado no alerta de concentração."),
        ("Limite de dedução do PGBL (% da renda bruta anual)", 0.12, FMT_PCT,
         "Vale só para quem declara no modelo completo e contribui para a Previdência Social."),
        ("IR sobre ganho de capital na venda de imóvel", 0.15, FMT_PCT,
         "Alíquota base. Há faixas maiores para ganhos altos e regras de isenção — ver aba Imóvel."),
        ("Dias por mês usados na projeção", 30, FMT_INT,
         "Faz as faixas de IR baterem exatamente (180, 360, 720 dias)."),
        ("Meses em que o come-cotas incide", "Maio e Novembro", FMT_TEXTO,
         "Último dia útil de maio e de novembro."),
        ("Meses de tolerância para renda vitalícia", 0, FMT_INT,
         "Reservado para uso futuro no cálculo de renda."),
    ]
    for i, (rotulo, val, fmt, obs) in enumerate(outros):
        lin = PRE_OUTROS_INI + i
        escreve(ws, f"A{lin}", rotulo, tam=9, borda=BORDA_GRADE)
        escreve(ws, f"B{lin}", val, tam=9, negrito=True, cor=COR_INPUT,
                bg=FILL_INPUT, fmt=fmt, alinha=AL_CENTRO, borda=BORDA_GRADE)
        escreve(ws, f"H{lin}", obs, tam=8, cor=CINZA_SUAVE, borda=BORDA_GRADE)

    ws.sheet_properties.tabColor = LARANJA
    link_aba(ws, "J4", AB_INICIO, "◀ INÍCIO")


# ==========================================================================
# ABA FONTES
# ==========================================================================
def construir_fontes(wb):
    ws = wb[AB_FONTES]
    esconde_grade(ws)
    titulo_pagina(ws, "FONTES DE DADOS E ATUALIZAÇÃO",
                  "De onde vem cada número, com que atraso, e como atualizar",
                  ate_col="G")
    larguras(ws, {"A": 34, "B": 32, "C": 52, "D": 42, "E": 34, "F": 18, "G": 46})

    aviso(ws, 4,
          "Nem todo dado da internet é 'tempo real'. A coluna PERIODICIDADE diz o atraso real de cada fonte. "
          "Um dado mensal com defasagem não pode ser apresentado ao cliente como cotação de hoje.",
          ate_col="G")

    cabecalho_tabela(ws, 6, 1,
                     ["O que alimenta", "Fonte", "Endereço", "Conteúdo",
                      "Periodicidade e atraso", "Atualização", "Como atualizar"],
                     altura=32)
    for i, f in enumerate(D.FONTES):
        lin = 7 + i
        for j, val in enumerate(f):
            col = get_column_letter(j + 1)
            escreve(ws, f"{col}{lin}", val, tam=8.5,
                    alinha=AL_ESQ_WRAP if j in (2, 3, 6) else AL_ESQ,
                    borda=BORDA_GRADE,
                    cor=AZUL if j == 5 else CINZA_TEXTO,
                    negrito=(j == 5))
        ws.row_dimensions[lin].height = 40

    lin = 7 + len(D.FONTES) + 1
    secao(ws, lin, "REGISTRO DE ATUALIZAÇÕES", ate_col="G")
    cabecalho_tabela(ws, lin + 1, 1,
                     ["Data", "Quem atualizou", "O que foi atualizado", "Fonte consultada",
                      "Observações", "", ""], altura=22)
    for k in range(12):
        for j in range(5):
            escreve(ws, f"{get_column_letter(j + 1)}{lin + 2 + k}", None, tam=9,
                    cor=COR_INPUT, bg=FILL_INPUT, borda=BORDA_GRADE)
        ws.row_dimensions[lin + 2 + k].height = 16

    ws.sheet_properties.tabColor = LARANJA
    link_aba(ws, "I1", AB_INICIO, "◀ INÍCIO")


# ==========================================================================
# ABA GLOSSÁRIO
# ==========================================================================
def construir_glossario(wb):
    ws = wb[AB_GLOSSARIO]
    esconde_grade(ws)
    titulo_pagina(ws, "EM PALAVRAS SIMPLES",
                  "Para explicar ao cliente sem usar jargão — leia em voz alta se ajudar",
                  ate_col="D")
    larguras(ws, {"A": 3, "B": 26, "C": 74, "D": 66})

    cabecalho_tabela(ws, 5, 2, ["Termo", "O que é, em uma frase", "Exemplo para usar na conversa"],
                     larguras=[26, 74, 66], altura=26)
    for i, (termo, definicao, exemplo) in enumerate(D.GLOSSARIO):
        lin = 6 + i
        bg = BRANCO if i % 2 == 0 else CINZA_FUNDO
        escreve(ws, f"B{lin}", termo, tam=10, negrito=True, cor=AZUL_ESCURO,
                bg=bg, alinha=AL_ESQ, borda=BORDA_GRADE)
        escreve(ws, f"C{lin}", definicao, tam=9.5, bg=bg, alinha=AL_ESQ_WRAP,
                borda=BORDA_GRADE)
        escreve(ws, f"D{lin}", exemplo, tam=9, cor=CINZA_SUAVE, bg=bg,
                italico=True, alinha=AL_ESQ_WRAP, borda=BORDA_GRADE)
        ws.row_dimensions[lin].height = 30

    ws.freeze_panes = "B6"
    ws.sheet_properties.tabColor = AZUL_MEDIO
    link_aba(ws, "D3", AB_INICIO, "◀ INÍCIO")
